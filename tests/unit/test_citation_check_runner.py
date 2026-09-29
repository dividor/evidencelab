"""Unit tests for the brief citation check runner (judge and DB faked)."""

import asyncio
import uuid
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from ui.backend.services import citation_check_runner as runner
from ui.backend.services.citation_fidelity import CitedPassage

pytestmark = pytest.mark.unit

EXCERPT = "Dropout fell from 12 to 7 percent in schools receiving meals."
SOURCE = {
    "index": 1,
    "title": "Kenya evaluation",
    "page": 4,
    "docId": "d1",
    "chunkId": "c1",
    "text": EXCERPT,
}
BRIEF = {
    "sections": [
        {
            "title": "Dropout",
            "status": "done",
            "content": "Dropout fell in fed schools [1]. Boys were unaffected [2].",
            "sources": [SOURCE],
        },
        {"title": "Draft", "status": "pending", "content": "", "sources": []},
    ]
}


def _passage(**kwargs) -> CitedPassage:
    base = dict(
        section_title="Dropout",
        passage="Dropout fell [1].",
        citation_indices=[1],
        sources=[SOURCE],
    )
    return CitedPassage(**{**base, **kwargs})


GOOD_REPLY = (
    '{"verdict": "supported", "confidence": 0.9, '
    '"supporting_quotes": [{"citation": 1, "quote": "Dropout fell from 12 to 7 percent"}], '
    '"problems": [], "explanation": "Matches."}'
)


class TestPassagesAndPrompts:
    def test_brief_passages_when_section_pending_then_only_done_sections(self):
        passages = runner.brief_passages(BRIEF)
        assert [p.passage for p in passages] == [
            "Dropout fell in fed schools [1].",
            "Boys were unaffected [2].",
        ]
        assert passages[1].missing_indices == [2]
        assert runner.count_cited_passages(BRIEF) == 2

    def test_judge_prompts_when_rendered_then_excerpt_and_quotes_unescaped(self):
        system, user = runner.judge_prompts(_passage(passage='Dropout "fell" [1].'))
        assert system.startswith("You are a meticulous fact-checker")
        assert 'Dropout "fell"' in user and "&quot;" not in user
        assert EXCERPT in user and '[1] document: "Kenya evaluation", p. 4' in user


class TestPassageRecord:
    def test_record_when_reply_valid_then_verdict_and_quote_verified(self):
        judgement = runner.judgement_or_error(_passage(), GOOD_REPLY)
        record = runner.passage_record(
            1, _passage(), judgement, {"prompt_tokens": 10, "completion_tokens": 5}
        )
        assert record["verdict"] == "supported" and record["flagged"] is False
        assert record["supporting_quotes"][0]["status"] == "verbatim"
        assert (
            record["quotes_verified"] == "1/1"
            and record["quote_not_in_source"] is False
        )
        assert record["sources"][0]["excerpt"] == EXCERPT
        assert record["prompt_tokens"] == 10 and record["error_message"] is None

    def test_record_when_quote_invented_then_flagged_not_in_source(self):
        reply = GOOD_REPLY.replace(
            "Dropout fell from 12 to 7 percent", "enrolment doubled overnight"
        )
        record = runner.passage_record(
            1, _passage(), runner.judgement_or_error(_passage(), reply)
        )
        assert record["quote_not_in_source"] is True
        assert record["supporting_quotes"][0]["status"] == "missing"

    def test_record_when_reply_malformed_then_cannot_assess_with_visible_error(self):
        judgement = runner.judgement_or_error(_passage(), "not json at all")
        record = runner.passage_record(
            1, _passage(), judgement, error=judgement["problems"][0]
        )
        assert record["verdict"] == "cannot_assess" and record["flagged"] is True
        assert record["error_message"].startswith("JUDGE ERROR")

    def test_record_when_dangling_citation_then_problem_added(self):
        passage = _passage(citation_indices=[1, 2], missing_indices=[2])
        judgement = runner.judgement_or_error(passage, GOOD_REPLY)
        record = runner.passage_record(1, passage, judgement)
        assert record["dangling_citations"] == "2"
        assert any("Dangling citation(s) [2]" in p for p in record["problems"])


class TestJudgePassage:
    @pytest.mark.asyncio
    async def test_judge_when_no_sources_then_not_sent_to_judge(self):
        calls = []

        async def judge(system, user):
            calls.append(user)
            return GOOD_REPLY, {}

        passage = _passage(sources=[], missing_indices=[1])
        record = await runner.judge_passage(1, passage, judge, asyncio.Semaphore(1))
        assert calls == []
        assert record["verdict"] == "cannot_assess"
        assert "No source stored for citation(s) [1]" in record["error_message"]

    @pytest.mark.asyncio
    async def test_judge_when_call_raises_then_error_row_not_exception(self):
        async def judge(system, user):
            raise RuntimeError("provider down")

        record = await runner.judge_passage(1, _passage(), judge, asyncio.Semaphore(1))
        assert record["verdict"] == "cannot_assess"
        assert record["error_message"] == "JUDGE ERROR: the model call failed"
        assert "provider down" not in record["error_message"]


class TestSummary:
    def _rows(self):
        return [
            {
                "brief_section": "A",
                "verdict": "supported",
                "flagged": False,
                "quote_not_in_source": False,
                "prompt_tokens": 10,
                "completion_tokens": 2,
            },
            {
                "brief_section": "A",
                "verdict": "unsupported",
                "flagged": True,
                "quote_not_in_source": True,
                "prompt_tokens": 10,
                "completion_tokens": 2,
            },
            {
                "brief_section": "B",
                "verdict": "partially_supported",
                "flagged": True,
                "quote_not_in_source": False,
                "prompt_tokens": 0,
                "completion_tokens": 0,
                "error_message": "JUDGE ERROR",
            },
        ]

    def test_summarise_when_rows_then_counts_sections_and_tokens(self, monkeypatch):
        monkeypatch.setattr(
            "ui.backend.utils.llm_costs.compute_cost", lambda m, p, c: 0.0123
        )
        stats = runner.summarise(self._rows(), 1500, "judge-x")
        assert stats["total"] == 3 and stats["flagged"] == 2
        assert stats["verdicts"] == {
            "supported": 1,
            "partially_supported": 1,
            "unsupported": 1,
            "cannot_assess": 0,
        }
        assert stats["flagged_share"] == 0.67
        assert stats["quote_not_in_source"] == 1 and stats["not_judged"] == 1
        assert [s["brief_section"] for s in stats["by_section"]] == ["A", "B"]
        assert stats["by_section"][0]["flagged_share"] == 0.5
        assert stats["total_tokens"] == 24 and stats["cost_usd"] == 0.0123

    def test_build_workbook_when_passages_then_three_sheets_flagged_by_severity(self):
        passages = [
            {
                "passage_id": 1,
                "brief_section": "A",
                "passage": "p1",
                "verdict": "partially_supported",
                "flagged": True,
                "confidence": 0.9,
                "problems": ["x"],
                "supporting_quotes": [{"citation": 1, "quote": "q", "status": "near"}],
            },
            {
                "passage_id": 2,
                "brief_section": "A",
                "passage": "p2",
                "verdict": "unsupported",
                "flagged": True,
                "confidence": 0.5,
                "problems": [],
                "supporting_quotes": [],
            },
            {
                "passage_id": 3,
                "brief_section": "A",
                "passage": "p3",
                "verdict": "supported",
                "flagged": False,
                "confidence": 1.0,
                "problems": [],
                "supporting_quotes": [],
            },
        ]
        book = load_workbook(
            BytesIO(
                runner.build_review_workbook(
                    passages, [{"brief_section": "A", "total": 3}]
                )
            )
        )
        assert book.sheetnames == ["Flagged", "All judgements", "Summary"]
        flagged = list(book["Flagged"].iter_rows(values_only=True))
        assert flagged[0][0] == "passage_id"
        assert [r[0] for r in flagged[1:]] == [2, 1]
        all_rows = list(book["All judgements"].iter_rows(values_only=True))
        assert len(all_rows) == 4 and "[1] q (near)" in all_rows[1]


class FakeSession:
    """Just enough AsyncSession for the runner: get/add/commit/execute."""

    def __init__(self, check, brief):
        self.rows = {("check", check.id): check, ("brief", brief.id): brief}
        self.added = []
        self.commits = 0

    async def get(self, model, key):
        kind = "brief" if model.__name__ == "Brief" else "check"
        return self.rows.get((kind, key))

    def add(self, row):
        self.added.append(row)

    async def commit(self):
        self.commits += 1

    async def execute(self, statement):
        return SimpleNamespace(rowcount=0)

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False


def _check_and_brief():
    brief = SimpleNamespace(id=uuid.uuid4(), content=BRIEF)
    check = SimpleNamespace(
        id=uuid.uuid4(),
        brief_id=brief.id,
        brief_title="B",
        data_source="wfp",
        created_by_user_id=uuid.uuid4(),
        judge_model="judge-x",
        status="pending",
        summary_stats=None,
        started_at=None,
        finished_at=None,
    )
    return check, brief


class TestRunCheck:
    @pytest.mark.asyncio
    async def test_run_when_judge_answers_then_rows_committed_and_completed(
        self, monkeypatch
    ):
        recorded = []

        async def fake_record(**kwargs):
            recorded.append(kwargs)

        monkeypatch.setattr(
            "ui.backend.services.usage_recorder.record_llm_usage", fake_record
        )
        monkeypatch.setattr(
            "ui.backend.utils.llm_costs.compute_cost", lambda m, p, c: None
        )
        check, brief = _check_and_brief()
        session = FakeSession(check, brief)

        async def judge(system, user):
            return GOOD_REPLY, {"prompt_tokens": 7, "completion_tokens": 3}

        await runner.run_check(check.id, session_factory=lambda: session, judge=judge)
        assert check.status == "completed"
        assert [
            r.passage_id for r in sorted(session.added, key=lambda r: r.passage_id)
        ] == [1, 2]
        by_id = {r.passage_id: r for r in session.added}
        assert by_id[1].verdict == "supported"
        assert by_id[2].verdict == "cannot_assess"  # no source stored for [2]
        stats = check.summary_stats
        assert (
            stats["total"] == 2
            and stats["flagged"] == 1
            and stats["prompt_tokens"] == 7
        )
        assert recorded and recorded[0]["activity_type"] == "evaluation"
        assert recorded[0]["search_id"] == check.id

    @pytest.mark.asyncio
    async def test_run_when_brief_missing_then_failed_with_message(self):
        check, brief = _check_and_brief()
        session = FakeSession(check, brief)
        del session.rows[("brief", brief.id)]
        await runner.run_check(check.id, session_factory=lambda: session)
        assert check.status == "failed"
        assert check.summary_stats == {"error": "Brief not found"}

    @pytest.mark.asyncio
    async def test_run_when_cancelled_midway_then_stops_without_completing(
        self, monkeypatch
    ):
        check, brief = _check_and_brief()
        session = FakeSession(check, brief)

        async def judge(system, user):
            check.status = "failed"  # the cancel route flips the row
            return GOOD_REPLY, {}

        await runner.run_check(check.id, session_factory=lambda: session, judge=judge)
        assert check.status == "failed"
        assert session.added == []

    @pytest.mark.asyncio
    async def test_run_when_execution_raises_then_generic_failure(self, monkeypatch):
        check, brief = _check_and_brief()
        session = FakeSession(check, brief)
        monkeypatch.setattr(runner, "brief_passages", lambda content: 1 / 0)
        await runner.run_check(check.id, session_factory=lambda: session)
        assert check.status == "failed"
        assert check.summary_stats == {"error": "Check failed"}


class TestResolveJudgeModel:
    def test_resolve_when_combo_has_summary_model_then_used(self, monkeypatch):
        monkeypatch.setattr(
            "ui.backend.services.test_runner._resolve_combo",
            lambda name: {"summary_model": "combo-model"} if name == "X" else {},
        )
        monkeypatch.setattr(
            "ui.backend.services.test_runner._default_summary_model", lambda: "default"
        )
        assert runner.resolve_judge_model("X") == "combo-model"
        assert runner.resolve_judge_model(None) == "default"
