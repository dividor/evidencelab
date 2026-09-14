"""Unit tests for notebooks/citation_fidelity_lib.py.

The module ports the frontend's citation parsing (CitedContent.tsx,
briefHighlights.ts) and adds the judge plumbing for the citation-fidelity
notebook; these tests pin both to the cases the notebook relies on.
"""

import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "notebooks"))

from citation_fidelity_lib import (  # noqa: E402
    FLAGGED_VERDICTS,
    QUOTE_MISSING,
    QUOTE_NEAR,
    QUOTE_VERBATIM,
    CitedPassage,
    build_judge_messages,
    extract_cited_numbers,
    extract_cited_passages,
    find_fragment,
    format_excerpts,
    judgement_row,
    locate_quote,
    match_key,
    parse_judge_response,
    parse_section_breadcrumb,
    passage_source_rows,
    quote_fragments,
    researched_sections,
    source_body,
    source_label,
    source_section,
    split_sentences,
    strip_citation_markers,
    verify_quotes,
    write_review_workbook,
)

SOURCE_1 = {
    "index": 1,
    "chunkId": "c1",
    "docId": "d1",
    "title": "Mali Evaluation",
    "page": 12,
    "headings": ["Findings", "3.2 Enrolment"],
    "text": "-- Findings > 3.2 Enrolment --\n\nEnrolment rose by 10 percentage points in Mali.",
}
SOURCE_2 = {
    "index": 2,
    "chunkId": "c2",
    "docId": "d2",
    "title": "Kenya Review",
    "text": "Numeracy improved in Kenya.",
}
SECTION = {
    "title": "Attendance",
    "status": "done",
    "content": (
        "## Attendance\n\nEnrolment rose by 10 points in Mali [1]. "
        "Numeracy improved in Kenya [2, 1]. No citation here. Dangling [9]."
    ),
    "sources": [SOURCE_1, SOURCE_2],
}


@pytest.mark.unit
class TestMarkdownParsing:
    def test_extract_cited_numbers_when_single_and_grouped_then_sorted_unique(self):
        assert extract_cited_numbers("One [3]. Two [1, 3]. Three [12].") == [1, 3, 12]

    def test_extract_cited_numbers_when_bracketed_words_then_ignored(self):
        assert extract_cited_numbers("See [appendix] and [1a].") == []

    def test_split_sentences_when_headings_and_lists_then_split_on_newlines(self):
        text = "## Title\n\nFirst sentence. Second one!\n- item [1]\n- item two"
        assert split_sentences(text) == [
            "## Title",
            "First sentence.",
            "Second one!",
            "- item [1]",
            "- item two",
        ]

    def test_strip_citation_markers_when_markers_then_removed_without_stray_space(self):
        assert strip_citation_markers("Rose by 10 points [1, 2]. Also [3] here.") == (
            "Rose by 10 points. Also here."
        )

    def test_parse_section_breadcrumb_when_leading_marker_line_then_split(self):
        section, body = parse_section_breadcrumb("-- A > B --\n\nBody text.")
        assert section == "A > B"
        assert body == "Body text."

    def test_parse_section_breadcrumb_when_no_marker_then_full_text(self):
        assert parse_section_breadcrumb("Plain text.") == (None, "Plain text.")


@pytest.mark.unit
class TestSourceHelpers:
    def test_source_section_when_headings_then_joined(self):
        assert source_section(SOURCE_1) == "Findings > 3.2 Enrolment"

    def test_source_section_when_no_headings_then_breadcrumb_fallback(self):
        src = {"text": "-- Summary > 2.1 --\nbody"}
        assert source_section(src) == "Summary > 2.1"

    def test_source_body_when_breadcrumb_then_stripped(self):
        assert (
            source_body(SOURCE_1) == "Enrolment rose by 10 percentage points in Mali."
        )

    def test_source_label_when_page_then_included(self):
        assert source_label(SOURCE_1) == "[1] Mali Evaluation, p. 12"
        assert source_label(SOURCE_2) == "[2] Kenya Review"

    def test_researched_sections_when_pending_or_empty_then_excluded(self):
        content = {
            "sections": [
                SECTION,
                {"title": "Empty", "status": "done", "content": " "},
                {"title": "Pending", "status": "pending", "content": "x"},
            ]
        }
        assert [s["title"] for s in researched_sections(content)] == ["Attendance"]


@pytest.mark.unit
class TestExtractCitedPassages:
    def test_extract_cited_passages_when_section_then_one_per_citing_sentence(self):
        passages = extract_cited_passages(SECTION)
        assert [p.citation_indices for p in passages] == [[1], [1, 2], [9]]
        assert passages[0].section_title == "Attendance"
        assert passages[0].passage_clean == "Enrolment rose by 10 points in Mali."

    def test_extract_cited_passages_when_multiple_citations_then_sources_in_index_order(
        self,
    ):
        passages = extract_cited_passages(SECTION)
        assert [s["chunkId"] for s in passages[1].sources] == ["c1", "c2"]
        assert not passages[1].dangling

    def test_extract_cited_passages_when_unknown_index_then_recorded_as_missing(self):
        dangling = extract_cited_passages(SECTION)[2]
        assert dangling.dangling
        assert dangling.missing_indices == [9]
        assert dangling.sources == []

    def test_passage_source_rows_when_passages_then_long_format_with_missing_marker(
        self,
    ):
        rows = passage_source_rows(extract_cited_passages(SECTION))
        assert [(r["passage_id"], r["citation"]) for r in rows] == [
            (1, 1),
            (2, 1),
            (2, 2),
            (3, 9),
        ]
        assert rows[0]["source_section"] == "Findings > 3.2 Enrolment"
        assert rows[0]["source_excerpt"].startswith("Enrolment rose")
        assert rows[3]["document"].startswith("(no source stored")


@pytest.mark.unit
class TestJudgePrompt:
    def test_format_excerpts_when_sources_then_numbered_with_section(self):
        passage = extract_cited_passages(SECTION)[1]
        text = format_excerpts(passage)
        assert (
            "### Excerpt [1] Mali Evaluation, p. 12 — section: Findings > 3.2 Enrolment"
            in text
        )
        assert "### Excerpt [2] Kenya Review\nNumeracy improved in Kenya." in text
        assert "-- Findings" not in text  # breadcrumb line stripped

    def test_build_judge_messages_when_template_then_fields_filled(self):
        passage = extract_cited_passages(SECTION)[1]
        messages = build_judge_messages(
            passage, "SYS", "S={brief_section}|P={passage}|C={citations}|E={excerpts}"
        )
        assert messages[0] == {"role": "system", "content": "SYS"}
        user = messages[1]["content"]
        assert user.startswith(
            "S=Attendance|P=Numeracy improved in Kenya.|C=[1], [2]|E="
        )
        assert "[2, 1]" not in user


@pytest.mark.unit
class TestJudgeResponse:
    def test_parse_judge_response_when_valid_json_then_normalised(self):
        raw = (
            '{"verdict": "Partially_Supported", "confidence": "0.8", '
            '"supporting_quotes": [{"citation": 1, "quote": " rose by 10 "}, {"quote": ""}], '
            '"problems": ["Kenya not in source", " "], "explanation": " Half right. "}'
        )
        parsed = parse_judge_response(raw)
        assert parsed["verdict"] == "partially_supported"
        assert parsed["confidence"] == 0.8
        assert parsed["supporting_quotes"] == [{"citation": 1, "quote": "rose by 10"}]
        assert parsed["problems"] == ["Kenya not in source"]
        assert parsed["explanation"] == "Half right."

    def test_parse_judge_response_when_fenced_then_unwrapped(self):
        parsed = parse_judge_response('```json\n{"verdict": "supported"}\n```')
        assert parsed["verdict"] == "supported"
        assert parsed["confidence"] == 0.0

    def test_parse_judge_response_when_unknown_verdict_then_raises(self):
        with pytest.raises(ValueError, match="Unknown verdict"):
            parse_judge_response('{"verdict": "maybe"}')

    def test_parse_judge_response_when_not_json_then_raises(self):
        with pytest.raises(ValueError, match="non-JSON"):
            parse_judge_response("Supported, I think.")

    def test_verify_quotes_when_in_cited_excerpt_then_status_per_quote(self):
        passage = extract_cited_passages(SECTION)[1]
        quotes = [
            {"citation": 1, "quote": "rose by 10 PERCENTAGE points"},
            {"citation": 2, "quote": "rose by 10 percentage points"},  # wrong excerpt
            {"citation": None, "quote": "numeracy improved"},  # any excerpt
            {"citation": 1, "quote": "invented sentence"},
        ]
        checks = verify_quotes(passage, quotes)
        assert [c["status"] for c in checks] == [
            QUOTE_VERBATIM,
            QUOTE_MISSING,
            QUOTE_VERBATIM,
            QUOTE_MISSING,
        ]
        assert checks[0]["spans"] == ["rose by 10 percentage points"]

    def test_judgement_row_when_unsupported_then_flagged_and_quote_check(self):
        passage = extract_cited_passages(SECTION)[0]
        row = judgement_row(
            7,
            passage,
            {
                "verdict": "unsupported",
                "confidence": 0.9,
                "supporting_quotes": [{"citation": 1, "quote": "not there"}],
                "problems": ["a", "b"],
                "explanation": "why",
            },
        )
        assert row["passage_id"] == 7
        assert row["flagged"] is True
        assert row["citations"] == "1"
        assert row["documents"] == "[1] Mali Evaluation"
        assert row["problems"] == "a\nb"
        assert row["quotes_verified"] == "0/1"
        assert row["quote_not_in_source"] is True

    def test_judgement_row_when_near_quote_then_counted_as_found(self):
        passage = extract_cited_passages(SECTION)[0]
        row = judgement_row(
            1,
            passage,
            {
                "verdict": "supported",
                "supporting_quotes": [
                    {"citation": 1, "quote": "rose by 10 percen tage points"}
                ],
            },
        )
        assert row["quotes_verified"] == "1/1 (1 near)"
        assert row["quote_not_in_source"] is False

    def test_judgement_row_when_supported_then_not_flagged(self):
        passage = extract_cited_passages(SECTION)[0]
        row = judgement_row(
            1, passage, {"verdict": "supported", "supporting_quotes": []}
        )
        assert row["flagged"] is False
        assert row["quotes_verified"] == "0/0"
        assert row["quote_not_in_source"] is False

    def test_flagged_verdicts_when_listed_then_exclude_supported(self):
        assert "supported" not in FLAGGED_VERDICTS


@pytest.mark.unit
class TestQuoteMatching:
    BODY = (
        "· In Kenya , at the Government's request,[^70] WFP launched an urban response "
        "in Nairobi's informal settlements and Mombasa. The G overnment agreed. "
        "Financed through the plan, WFP pre-positioned food."
    )

    def test_match_key_when_punctuation_footnotes_unicode_then_normalised(self):
        assert match_key("Government’s  “drive-through”,[^7] café") == (
            "government s drive through café"
        )

    def test_quote_fragments_when_ellipsis_then_split(self):
        assert quote_fragments("first part ... second part … third [...] fourth") == [
            "first part",
            "second part",
            "third",
            "fourth",
        ]
        assert quote_fragments("...") == []

    def test_find_fragment_when_punctuation_differs_then_verbatim_with_span(self):
        status, span = find_fragment(
            "In Kenya, at the Government's request, WFP launched", self.BODY
        )
        assert status == QUOTE_VERBATIM
        assert span == "In Kenya , at the Government's request,[^70] WFP launched"

    def test_find_fragment_when_pdf_split_word_then_near(self):
        status, span = find_fragment("The Government agreed.", self.BODY)
        assert status == QUOTE_NEAR
        assert span == "The G overnment agreed."

    def test_find_fragment_when_not_in_body_then_missing(self):
        assert find_fragment("rations were doubled", self.BODY) == (QUOTE_MISSING, None)
        assert find_fragment("", self.BODY) == (QUOTE_MISSING, None)

    def test_locate_quote_when_fragments_across_bodies_then_worst_status(self):
        result = locate_quote(
            "WFP launched an urban response ... pre-positioned food", [self.BODY]
        )
        assert result["status"] == QUOTE_VERBATIM
        assert result["spans"] == [
            "WFP launched an urban response",
            "pre-positioned food.",
        ]
        assert locate_quote("The Government agreed ... nothing here", [self.BODY])[
            "status"
        ] == (QUOTE_MISSING)
        assert locate_quote("The Government agreed ... urban response", [self.BODY])[
            "status"
        ] == (QUOTE_NEAR)
        assert locate_quote("urban response", ["other text", self.BODY])["status"] == (
            QUOTE_VERBATIM
        )


@pytest.mark.unit
class TestWorkbook:
    def test_write_review_workbook_when_sheets_then_readable_back(self, tmp_path):
        path = tmp_path / "review.xlsx"
        sheets = {
            "Citations": [
                {"passage": "a", "page": 1, "problems": ["x", "y"], "meta": {"k": 1}},
                {"passage": "b", "page": None, "problems": [], "meta": {}},
            ],
            "A very long sheet name that exceeds the limit": [],
        }
        write_review_workbook(sheets, str(path))
        book = load_workbook(path)
        assert book.sheetnames == ["Citations", "A very long sheet name that exc"]
        sheet = book["Citations"]
        assert [c.value for c in sheet[1]] == ["passage", "page", "problems", "meta"]
        assert [c.value for c in sheet[2]] == ["a", 1, "x\ny", '{"k": 1}']
        assert sheet.freeze_panes == "A2"
        assert sheet.column_dimensions["A"].width == 60  # wide-column preset

    def test_write_review_workbook_when_nan_then_empty_cell(self, tmp_path):
        path = tmp_path / "nan.xlsx"
        write_review_workbook({"S": [{"v": float("nan")}]}, str(path))
        assert load_workbook(path)["S"]["A2"].value is None


@pytest.mark.unit
def test_cited_passage_dataclass_when_no_sources_then_clean_text_still_available():
    passage = CitedPassage(section_title="S", passage="Fact [4].", citation_indices=[4])
    assert passage.passage_clean == "Fact."
    assert not passage.dangling
