"""Helpers for the "LLM judge citation fidelity" notebook.

The notebook keeps everything a reviewer wants to *see* (prompts, verdicts,
tables) in its cells; this module holds the plumbing that is better unit
tested than eyeballed: splitting a brief section into cited passages, pairing
each passage with the exact source excerpts it cites, building the judge
prompt, validating the judge's answer, and writing the review spreadsheets.

The citation parsing mirrors the frontend so the notebook sees a stored brief
exactly the way the UI renders it:

- ``ui/frontend/src/components/citations/CitedContent.tsx`` (``CITATION_REGEX``)
- ``ui/frontend/src/components/brief/briefHighlights.ts`` (``SENTENCE_SPLIT_RE``)

Tests: ``tests/unit/test_citation_fidelity_lib.py``.
"""

import json
import re
import unicodedata
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

# `[1]`, `[1, 3]` — same as CITATION_REGEX in CitedContent.tsx.
CITATION_RE = re.compile(r"\[(\d+(?:,\s*\d+)*)\]")

# Sentences end at ./!/? followed by whitespace; newlines (headings, list
# items) break sentences too — same as SENTENCE_SPLIT_RE in briefHighlights.ts.
SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+|\n+")

# A leading "-- h1 > h2 -- " line is the chunk's heading breadcrumb.
SECTION_BREADCRUMB_RE = re.compile(r"^\s*--\s*(.+?)\s*--\s*$")

# The only verdicts the judge may return; anything else is a parse error.
JUDGE_VERDICTS = ("supported", "partially_supported", "unsupported", "cannot_assess")

# Verdicts that put a passage on the human review list.
FLAGGED_VERDICTS = ("partially_supported", "unsupported", "cannot_assess")


# ---------------------------------------------------------------------------
# Brief markdown → cited passages
# ---------------------------------------------------------------------------


def parse_citation_numbers(raw: str) -> List[int]:
    """``"1, 3"`` -> ``[1, 3]``."""
    return [int(part) for part in raw.split(",") if part.strip().isdigit()]


def extract_cited_numbers(text: str) -> List[int]:
    """All unique citation indices appearing in the markdown, sorted."""
    cited: set = set()
    for match in CITATION_RE.finditer(text):
        cited.update(parse_citation_numbers(match.group(1)))
    return sorted(cited)


def split_sentences(markdown: str) -> List[str]:
    """Section markdown -> claim-sized sentences (same splitter as the UI)."""
    return [s.strip() for s in SENTENCE_SPLIT_RE.split(markdown) if s.strip()]


def strip_citation_markers(text: str) -> str:
    """Remove ``[n]`` / ``[n, m]`` markers, collapsing the space they leave."""
    stripped = CITATION_RE.sub("", text)
    stripped = re.sub(r"[ \t]+([.,;:!?])", r"\1", stripped)
    return re.sub(r"[ \t]{2,}", " ", stripped).strip()


def parse_section_breadcrumb(text: str) -> Tuple[Optional[str], str]:
    """Split a stored excerpt into (heading breadcrumb, body)."""
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines) and lines[i].strip() == "":
        i += 1
    match = SECTION_BREADCRUMB_RE.match(lines[i].strip()) if i < len(lines) else None
    if match:
        body = "\n".join(lines[i + 1 :]).lstrip("\n")
        return match.group(1).strip(), body
    return None, text


def researched_sections(content: Dict[str, Any]) -> List[Dict[str, Any]]:
    """The brief sections that were actually researched (have text)."""
    return [
        s
        for s in content.get("sections") or []
        if s.get("status") == "done" and (s.get("content") or "").strip()
    ]


def source_section(source: Dict[str, Any]) -> str:
    """Where in the source document the excerpt sits (its heading path)."""
    headings = [h for h in source.get("headings") or [] if h]
    if headings:
        return " > ".join(headings)
    breadcrumb, _ = parse_section_breadcrumb(source.get("text") or "")
    return breadcrumb or ""


def source_body(source: Dict[str, Any]) -> str:
    """The excerpt text without its leading heading-breadcrumb line."""
    _, body = parse_section_breadcrumb(source.get("text") or "")
    return body.strip()


def source_label(source: Dict[str, Any]) -> str:
    """Short human label: ``[3] Title, p. 70``."""
    page = source.get("page")
    page_part = f", p. {page}" if page is not None else ""
    return f"[{source.get('index')}] {source.get('title') or 'Untitled'}{page_part}"


@dataclass
class CitedPassage:
    """One brief sentence carrying ``[n]`` markers, with the sources it cites."""

    section_title: str
    passage: str
    citation_indices: List[int]
    sources: List[Dict[str, Any]] = field(default_factory=list)
    missing_indices: List[int] = field(default_factory=list)

    @property
    def passage_clean(self) -> str:
        return strip_citation_markers(self.passage)

    @property
    def dangling(self) -> bool:
        """Cites a number that has no source in the section."""
        return bool(self.missing_indices)


def extract_cited_passages(section: Dict[str, Any]) -> List[CitedPassage]:
    """Every sentence of a section that cites something, with its sources.

    Sources are matched by the ``index`` the markdown refers to. A number
    with no matching source is recorded in ``missing_indices`` rather than
    silently dropped, so dangling citations surface in the review.
    """
    sources_by_index = {
        src["index"]: src
        for src in section.get("sources") or []
        if src.get("index") is not None
    }
    passages: List[CitedPassage] = []
    for sentence in split_sentences(section.get("content") or ""):
        indices = extract_cited_numbers(sentence)
        if not indices:
            continue
        passage = CitedPassage(
            section_title=section.get("title", ""),
            passage=sentence,
            citation_indices=indices,
        )
        for index in indices:
            source = sources_by_index.get(index)
            if source is None:
                passage.missing_indices.append(index)
            else:
                passage.sources.append(source)
        passages.append(passage)
    return passages


def passage_source_rows(passages: Iterable[CitedPassage]) -> List[Dict[str, Any]]:
    """Long table: one row per (passage, cited source) — the Excel export."""
    rows: List[Dict[str, Any]] = []
    for pid, passage in enumerate(passages, start=1):
        for source in passage.sources:
            rows.append(
                {
                    "passage_id": pid,
                    "brief_section": passage.section_title,
                    "passage": passage.passage_clean,
                    "citation": source.get("index"),
                    "document": source.get("title"),
                    "page": source.get("page"),
                    "source_section": source_section(source),
                    "source_excerpt": source_body(source),
                    "doc_id": source.get("docId"),
                    "chunk_id": source.get("chunkId"),
                    "pdf_url": source.get("pdfUrl"),
                }
            )
        for index in passage.missing_indices:
            rows.append(
                {
                    "passage_id": pid,
                    "brief_section": passage.section_title,
                    "passage": passage.passage_clean,
                    "citation": index,
                    "document": "(no source stored for this citation number)",
                }
            )
    return rows


# ---------------------------------------------------------------------------
# Judge prompt and response
# ---------------------------------------------------------------------------


def format_excerpts(passage: CitedPassage) -> str:
    """The cited excerpts as the judge sees them, numbered by citation."""
    blocks = []
    for source in passage.sources:
        header = source_label(source)
        section = source_section(source)
        if section:
            header += f" — section: {section}"
        blocks.append(f"### Excerpt {header}\n{source_body(source)}")
    return "\n\n".join(blocks)


def build_judge_messages(
    passage: CitedPassage, system_prompt: str, user_template: str
) -> List[Dict[str, str]]:
    """Chat messages for one passage. ``user_template`` is ``str.format``-ed
    with ``brief_section``, ``passage``, ``citations`` and ``excerpts``."""
    user = user_template.format(
        brief_section=passage.section_title,
        passage=passage.passage_clean,
        citations=", ".join(f"[{i}]" for i in passage.citation_indices),
        excerpts=format_excerpts(passage),
    )
    return [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user},
    ]


# Footnote markers the parser leaves in chunk text, e.g. "[^70]".
FOOTNOTE_RE = re.compile(r"\[\^\d+\]")

# How a judge marks a gap between two quoted fragments: an ellipsis, or a
# blank line when it copies two separate paragraphs or list items.
ELLIPSIS_RE = re.compile(r"\s*(?:\[?\.\.\.\]?|\[?…\]?)\s*|\n[ \t]*\n\s*")

# Similarity (0-1) above which a quote counts as "near-verbatim": the same
# words modulo PDF extraction artefacts such as "G overnment" or "highl y".
NEAR_MATCH_THRESHOLD = 0.9

# Quote check outcomes, best to worst.
QUOTE_VERBATIM = "verbatim"
QUOTE_NEAR = "near"
QUOTE_MISSING = "missing"


def match_key(text: str) -> str:
    """Comparison form: unicode-normalised, lower-case, footnote markers and
    punctuation dropped, whitespace collapsed. Punctuation differences (curly
    quotes, bullets, dashes) are never what a reviewer cares about."""
    text = unicodedata.normalize("NFKC", FOOTNOTE_RE.sub(" ", text)).lower()
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def quote_fragments(quote: str) -> List[str]:
    """A quote may join distant sentences with "..." or a blank line; each
    piece is checked on its own."""
    return [part for part in ELLIPSIS_RE.split(quote) if match_key(part)]


def _keyed_words(text: str) -> List[Tuple[str, int]]:
    """(comparison token, index of the original whitespace-split word)."""
    keyed: List[Tuple[str, int]] = []
    for index, word in enumerate(text.split()):
        keyed.extend((token, index) for token in match_key(word).split())
    return keyed


def _window_ratio(window: List[str], needle: List[str]) -> float:
    if window == needle:
        return 1.0
    return SequenceMatcher(None, " ".join(window), " ".join(needle)).ratio()


def find_fragment(fragment: str, body: str) -> Tuple[str, Optional[str]]:
    """Locate one quoted fragment in an excerpt.

    Returns ``(status, span)``: ``verbatim`` when the words match exactly
    (punctuation aside), ``near`` when a window of the excerpt is at least
    ``NEAR_MATCH_THRESHOLD`` similar, else ``missing``. ``span`` is the
    excerpt text that matched, for highlighting.
    """
    needle = match_key(fragment).split()
    keyed = _keyed_words(body)
    words = body.split()
    if not needle or not keyed:
        return QUOTE_MISSING, None
    anchors = set(needle[:3])
    tokens = [token for token, _ in keyed]
    best_ratio, best_span = 0.0, None
    for at, token in enumerate(tokens):
        if token not in anchors:
            continue
        for start in range(max(0, at - 2), at + 1):
            for width in range(max(1, len(needle) - 1), len(needle) + 4):
                ratio = _window_ratio(tokens[start : start + width], needle)
                if ratio > best_ratio:
                    first, last = (
                        keyed[start][1],
                        keyed[min(start + width, len(keyed)) - 1][1],
                    )
                    best_ratio, best_span = ratio, " ".join(words[first : last + 1])
                if ratio == 1.0:
                    return QUOTE_VERBATIM, best_span
    if best_ratio >= NEAR_MATCH_THRESHOLD:
        return QUOTE_NEAR, best_span
    return QUOTE_MISSING, None


def locate_quote(quote: str, bodies: Iterable[str]) -> Dict[str, Any]:
    """Check a whole quote against one or more excerpts.

    ``status`` is the worst fragment outcome (all verbatim → verbatim; any
    near → near; any missing → missing); ``spans`` are the excerpt passages
    that matched, for highlighting.
    """
    fragments = quote_fragments(quote)
    if not fragments:
        return {"status": QUOTE_MISSING, "spans": []}
    statuses: List[str] = []
    spans: List[str] = []
    for fragment in fragments:
        status, span = QUOTE_MISSING, None
        for body in bodies:
            status, span = find_fragment(fragment, body)
            if status != QUOTE_MISSING:
                break
        statuses.append(status)
        if span:
            spans.append(span)
    if QUOTE_MISSING in statuses:
        return {"status": QUOTE_MISSING, "spans": spans}
    if QUOTE_NEAR in statuses:
        return {"status": QUOTE_NEAR, "spans": spans}
    return {"status": QUOTE_VERBATIM, "spans": spans}


def verify_quotes(
    passage: CitedPassage, quotes: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """For each judge quote: was it found in the excerpt it is attributed to
    (or, without a citation number, in any cited excerpt)?

    The judge is asked to quote the source verbatim; checking that it did is
    the cheapest guard against a verdict resting on invented evidence. Each
    result is ``{"status": verbatim|near|missing, "spans": [...]}``.
    """
    bodies = {src.get("index"): source_body(src) for src in passage.sources}
    results = []
    for quote in quotes:
        citation = quote.get("citation")
        candidates = [bodies[citation]] if citation in bodies else list(bodies.values())
        results.append(locate_quote(quote["quote"], candidates))
    return results


def _clean_quotes(raw_quotes: Any) -> List[Dict[str, Any]]:
    quotes: List[Dict[str, Any]] = []
    for item in raw_quotes or []:
        if not isinstance(item, dict) or not item.get("quote"):
            continue
        citation_text = str(item.get("citation") or "")
        quotes.append(
            {
                "citation": int(citation_text) if citation_text.isdigit() else None,
                "quote": str(item["quote"]).strip(),
            }
        )
    return quotes


def parse_judge_response(raw: str) -> Dict[str, Any]:
    """Validate the judge's JSON answer into a fixed shape.

    Raises ``ValueError`` on malformed output so a bad response is visible
    instead of silently counted as a verdict.
    """
    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Judge returned non-JSON output: {raw[:200]!r}") from exc
    verdict = str(data.get("verdict", "")).strip().lower()
    if verdict not in JUDGE_VERDICTS:
        raise ValueError(
            f"Unknown verdict {verdict!r}; expected one of {JUDGE_VERDICTS}"
        )
    try:
        confidence = float(data.get("confidence", 0.0))
    except (TypeError, ValueError):
        confidence = 0.0
    return {
        "verdict": verdict,
        "confidence": max(0.0, min(1.0, confidence)),
        "supporting_quotes": _clean_quotes(data.get("supporting_quotes")),
        "problems": [str(p) for p in data.get("problems") or [] if str(p).strip()],
        "explanation": str(data.get("explanation") or "").strip(),
    }


def judgement_row(
    pid: int, passage: CitedPassage, judgement: Dict[str, Any]
) -> Dict[str, Any]:
    """One flat row per judged passage for the tables and the spreadsheet."""
    quotes = judgement.get("supporting_quotes") or []
    statuses = [check["status"] for check in verify_quotes(passage, quotes)]
    found = sum(status != QUOTE_MISSING for status in statuses)
    near = sum(status == QUOTE_NEAR for status in statuses)
    verdict = judgement["verdict"]
    return {
        "passage_id": pid,
        "brief_section": passage.section_title,
        "passage": passage.passage_clean,
        "citations": ", ".join(str(i) for i in passage.citation_indices),
        "documents": " | ".join(
            f"[{src.get('index')}] {src.get('title')}" for src in passage.sources
        ),
        "verdict": verdict,
        "flagged": verdict in FLAGGED_VERDICTS,
        "confidence": judgement.get("confidence"),
        "problems": "\n".join(judgement.get("problems") or []),
        "explanation": judgement.get("explanation", ""),
        "supporting_quotes": "\n".join(
            f"[{q.get('citation')}] {q['quote']}" for q in quotes
        ),
        "quotes_verified": f"{found}/{len(statuses)}"
        + (f" ({near} near)" if near else ""),
        "quote_not_in_source": QUOTE_MISSING in statuses,
    }


# ---------------------------------------------------------------------------
# Spreadsheet output
# ---------------------------------------------------------------------------

# Excel column widths (characters) for the long-text columns; others autosize.
_WIDE_COLUMNS = {
    "passage": 60,
    "source_excerpt": 80,
    "explanation": 60,
    "problems": 50,
    "supporting_quotes": 60,
    "documents": 45,
    "document": 45,
    "source_section": 35,
    "brief_section": 28,
}


def write_review_workbook(sheets: Dict[str, List[Dict[str, Any]]], path: str) -> None:
    """Write row dicts (e.g. ``DataFrame.to_dict("records")``) to one .xlsx,
    one sheet per key, with wrapped text, sensible widths and a frozen header
    row, so a reviewer can read it without reformatting.

    Column order follows the first row of each sheet. Uses openpyxl directly
    (no pandas) so the module stays importable in the app's test environment.
    """
    book = Workbook()
    book.remove(book.worksheets[0])
    for name, rows in sheets.items():
        sheet = book.create_sheet(title=name[:31])
        columns = list(rows[0].keys()) if rows else []
        sheet.append(columns)
        for row in rows:
            sheet.append([_cell_value(row.get(column)) for column in columns])
        sheet.freeze_panes = "A2"
        for col_idx, column in enumerate(columns, start=1):
            header = sheet.cell(row=1, column=col_idx)
            header.font = Font(bold=True)
            width = _WIDE_COLUMNS.get(str(column))
            if width is None:
                longest = max(
                    [len(str(column))] + [len(str(r.get(column, ""))) for r in rows]
                )
                width = min(max(longest + 2, 8), 40)
            sheet.column_dimensions[header.column_letter].width = width
            for row_idx in range(2, len(rows) + 2):
                sheet.cell(row=row_idx, column=col_idx).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
    book.save(path)


def _cell_value(value: Any) -> Any:
    """Excel cells take scalars only; lists/dicts become readable text."""
    if isinstance(value, (list, tuple)):
        return "\n".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    if value != value:  # NaN from pandas
        return None
    return value
